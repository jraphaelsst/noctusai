"""Contract typography (`contrato_gerador/estilo.py`) — the emphasis measured
from the office's signed contracts, asserted on the RENDERED `.docx`.

Synthetic fixtures only (`contrato_gerador_fixtures`). Rendering uses the
REAL seed docxtpl adapter; emphasis rides the matrícula extractor's own
`**`/`<u>` markup model and must never reach the document as a literal.
"""
from __future__ import annotations

import io
from dataclasses import replace

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from noctusai_lib.integrations.documents import has_raw_markup
from noctusai_lib.integrations.documents.formatting import FormatRange
from noctusai_lib.integrations.docx_render import get_docx_render_adapter

from app.modules.card_hub.contrato_gerador import derivacao, documento, estilo, lint
from tests.modules.card_hub import contrato_gerador_fixtures as fx


def _render(n: int, d=None):
    d = d if d is not None else fx.variante(n)
    pol = fx.politica_variante(n)
    hoje = derivacao._hoje_padrao()
    sw = derivacao.derivar_switches(d, pol, hoje)
    av = derivacao.avaliar(d, sw, pol, fx.ASSINATURA, hoje)
    assert av.pronto, (av.faltando, av.bloqueios)
    return documento.renderizar(get_docx_render_adapter(real=True), d, sw, pol, fx.ASSINATURA, hoje)


def _doc(r):
    return Document(io.BytesIO(r.docx))


def _paragrafo(doc, comeco: str):
    return next(p for p in doc.paragraphs if p.text.startswith(comeco))


def _runs(p) -> list[tuple[str, bool, bool]]:
    return [(r.text, bool(r.bold), bool(r.underline)) for r in p.runs if r.text]


def _negritos(p) -> list[str]:
    return [t for t, b, _u in _runs(p) if b]


@pytest.fixture(scope="module")
def r1():
    return _render(1)


@pytest.fixture(scope="module")
def doc1(r1):
    return _doc(r1)


class TestNoMarkerReachesTheDocument:
    @pytest.mark.parametrize("n", range(1, 7))
    @pytest.mark.parametrize("modalidade", ["digital", "fisica"])
    def test_every_variant_renders_clean(self, n, modalidade):
        r = _render(n, replace(fx.variante(n), modalidade_assinatura=modalidade))
        xml = documento.document_xml(r.docx)
        assert "**" not in xml
        assert "&lt;u&gt;" not in xml and "&lt;/u&gt;" not in xml
        assert not any(has_raw_markup(p) for p in r.paragrafos)
        assert lint.lint(r.paragrafos, referencias=r.referencias, clausulas=r.clausulas) == []

    def test_the_lint_refuses_an_unconverted_marker(self):
        hits = lint.lint(["O **VENDEDOR sem par."], referencias={}, clausulas={})
        assert "MARCACAO_NAO_CONVERTIDA" in {h["codigo"] for h in hits}

    def test_two_renders_are_identical(self, r1):
        assert documento.document_xml(_render(1).docx) == documento.document_xml(r1.docx)


class TestHeadingsAndTitle:
    def test_clause_label_underlined_topic_not(self, doc1):
        p = _paragrafo(doc1, "CLÁUSULA PRIMEIRA")
        assert p.style.name == "Heading 1"
        assert _runs(p) == [
            ("CLÁUSULA PRIMEIRA", False, True),
            (" – DO OBJETO DO CONTRATO", False, False),
        ]

    def test_every_clause_label_is_underlined(self, doc1):
        headings = [p for p in doc1.paragraphs if p.text.startswith("CLÁUSULA ")]
        assert len(headings) >= 10
        for p in headings:
            primeiro = p.runs[0]
            assert primeiro.text.startswith("CLÁUSULA ") and primeiro.underline

    def test_heading_and_title_styles_are_bold_and_black_arial(self, doc1):
        for nome in ("Title", "Heading 1"):
            s = doc1.styles[nome]
            assert s.font.bold is True
            assert s.font.name == estilo.FONTE
            assert s.font.size.pt == estilo.TAMANHO_PT
            assert str(s.font.color.rgb) == "000000"
        assert doc1.styles["Title"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert doc1.styles["Heading 1"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert doc1.styles["Title"].element.pPr.find(qn("w:pBdr")) is None

    def test_body_is_justified_arial_11(self, doc1):
        normal = doc1.styles["Normal"]
        assert normal.font.name == estilo.FONTE and normal.font.size.pt == estilo.TAMANHO_PT
        assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_the_quote_style_is_not_italic_nor_indented(self, doc1):
        q = doc1.styles["Quote"]
        assert q.font.italic is False
        assert q.paragraph_format.left_indent == 0


class TestQualificacaoAndTerms:
    def test_party_name_bold_documents_plain(self, doc1):
        p = _paragrafo(doc1, "Pelo presente")
        negritos = _negritos(p)
        assert "FULANO DE TAL" in negritos
        assert '"VENDEDOR"' in negritos  # quotes included, like the corpus
        assert not any("123.456.789-09" in t or "RG" in t for t in negritos)

    def test_defined_terms_bold_articles_plain(self, doc1):
        p = _paragrafo(doc1, "O VENDEDOR compromete-se")
        runs = _runs(p)
        assert ("VENDEDOR", True, False) in runs and ("COMPRADORA", True, False) in runs
        assert ("O ", False, False) in runs

    def test_uppercase_parte_is_bold(self, doc1):
        p = _paragrafo(doc1, "Parágrafo Segundo: Fica ajustado que em caso de rescisão")
        assert "PARTE" in _negritos(p)


class TestMoneyAndLabels:
    def test_price_with_its_extenso_is_bold(self, doc1):
        p = _paragrafo(doc1, "O VENDEDOR compromete-se")
        assert any(t.startswith("R$ ") and t.endswith("reais)") for t in _negritos(p))

    def test_parcela_label_sinal_and_value_bold(self, doc1):
        p = _paragrafo(doc1, "Parcela 01:")
        negritos = _negritos(p)
        assert negritos[:3] == ["Parcela 01:", "Sinal e princípio de pagamento:", "R$ 50.000,00 (cinquenta mil reais)"]
        # the favorecido: the VENDEDOR term, and his name upper-case + bold
        assert "FULANO DE TAL" in negritos

    def test_paragrafo_label_bold_colon_included(self, doc1):
        p = _paragrafo(doc1, "Parágrafo Único:")
        assert _runs(p)[0] == ("Parágrafo Único:", True, False)

    def test_prazos_and_percentages_stay_plain(self, doc1):
        todos = [t for p in doc1.paragraphs for t in _negritos(p)]
        assert not any("dias" in t or "%" in t for t in todos)


class TestMatriculaQuote:
    def test_label_inscricao_and_matricula_digits_bold(self, doc1):
        p = _paragrafo(doc1, "IMÓVEL:")
        assert p.style.name == "Quote"
        negritos = _negritos(p)
        assert negritos[0] == "IMÓVEL:"
        assert "12.345" in negritos  # matrícula digits
        assert not any("Matrícula Nº" in t for t in negritos)
        assert not any("Registro" in t or "Cartório" in t for t in negritos)

    def test_the_captured_formatting_survives_next_to_the_generators(self):
        """The quote keeps the bold/underline captured at transcription
        (`Matricula.formatacao`), alongside the generator's own emphasis."""
        texto = fx.MATRICULA_TEXTO
        ini_b, ini_u = texto.index("FULANO DE TAL"), texto.index("R.1/12.345")
        d = fx.variante(1)
        d = replace(
            d,
            matricula=replace(
                d.matricula,
                formatacao=(
                    FormatRange(start=ini_b, end=ini_b + len("FULANO DE TAL"), bold=True),
                    FormatRange(start=ini_u, end=ini_u + len("R.1/12.345"), underline=True),
                ),
            ),
        )
        p = _paragrafo(_doc(_render(1, d)), "IMÓVEL:")
        runs = _runs(p)
        assert ("FULANO DE TAL", True, False) in runs
        assert ("R.1/12.345", False, True) in runs
        assert runs[0] == ("IMÓVEL:", True, False)


class TestCertidoesAndSignatures:
    def test_group_header_bold_items_plain(self, doc1):
        cabecalho = _paragrafo(doc1, "1 - Em nome de")
        assert all(b for _t, b, _u in _runs(cabecalho))
        item = _paragrafo(doc1, "1.1 – ")
        assert _negritos(item) == []
        imovel = next(p for p in doc1.paragraphs if "Em Relação ao Imóvel" in p.text)
        assert all(b for _t, b, _u in _runs(imovel))

    def test_pendencia_letter_bold_text_plain(self, doc1):
        p = _paragrafo(doc1, "a-) ")
        assert _runs(p)[0] == ("a-)", True, False)
        assert all(not b for _t, b, _u in _runs(p)[1:])

    def test_date_line_right_aligned_and_plain(self, doc1):
        p = next(p for p in doc1.paragraphs if p.text.endswith("de 2026.") and p.text.startswith("Cidade Exemplo"))
        assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert _negritos(p) == []

    def test_signature_block_names_and_documents_bold(self):
        doc = _doc(_render(1, replace(fx.variante(1), modalidade_assinatura="fisica")))
        textos = [p.text for p in doc.paragraphs]
        corte = textos.index(next(t for t in textos if t.startswith("Cidade Exemplo,")))
        for p in doc.paragraphs[corte + 1:]:
            if p.text.startswith("___") or p.style.name == "Heading 1":
                assert _negritos(p) == []  # the rule line; TESTEMUNHAS: is bold by style
                continue
            assert all(b for _t, b, _u in _runs(p)), p.text


class TestEstiloHelpers:
    def test_negrito_and_sublinhado(self):
        assert estilo.negrito("X") == "**X**"
        assert estilo.sublinhado("X") == "<u>X</u>"
        assert estilo.negrito("") == "" and estilo.sublinhado("") == ""

    def test_nome_parte_upper_and_bold(self):
        assert estilo.nome_parte(" Fulana de Tal ") == "**FULANA DE TAL**"

    @pytest.mark.parametrize("marcado", ["já **negrito**", "meio <u>sub", "fim</u>"])
    def test_a_value_already_carrying_a_marker_is_refused(self, marcado):
        with pytest.raises(ValueError):
            estilo.negrito(marcado)

    def test_texto_plano_uses_the_same_parse_markup(self):
        assert estilo.texto_plano("<u>CLÁUSULA X</u> – **A**") == "CLÁUSULA X – A"
