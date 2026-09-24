"""Build the template .docx, render it through the seed `docx_render` adapter,
read the result back — then (`gerar_pdf`) turn that `.docx` into the ABNT
PDF the user actually receives (contract `projects/abnt-formatting-CONTRACT.md`
§5). Both encodings of the SAME rendering are now stored versions —
`contratos_service.nova_versao_gerada` saves the rendered `.docx` alongside
the PDF this module derives from it (2026-09-16, roadmap
`social-wiring-contract-automation-2026-09.md` question Q-artifact), so the
office can mark up the editable file without a second generation. This
module itself is unaware of storage either way — it only renders and
converts; see `contratos_service.py` for what happens to the bytes.

- The template is BUILT from `modelo_texto` with python-docx at runtime and
  cached — the wording stays reviewable text in git, no binary blob.
- Rendering is two passes over the same template: pass 1 only counts each
  clause's `par()` calls, pass 2 labels them ("Único" when a clause rendered
  exactly one). The totals are measured from the template itself, so there
  is no second, hand-kept paragraph declaration to drift from it.
- The product never imports docxtpl: `adapter` is the seed's
  `DocxRenderAdapter` (StrictUndefined — a missing key is an error; its
  `rich_text()` is likewise the only sanctioned way to build a docxtpl
  `{{r ... }}` context value — see contract §5's matrícula quote).
- ABNT by construction (§5.2): `template_bytes()` carries no font/bolding
  of its own — every paragraph gets a docx PARAGRAPH STYLE
  (`Title`/`Heading 1`/`Quote`/default `Normal`), and `gerar_pdf` maps that
  STYLE (never the source text) to a `ParagraphKind` before handing it to
  the seed's `render_abnt_pdf` — the PDF's margins/fonts/spacing come
  entirely from the KIND, never from this module's own formatting choices.
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from typing import Optional

from docx import Document

from noctusai_lib.integrations.docx_render import DocxRenderAdapter
from noctusai_lib.integrations.documents.abnt import paragraphs_from_docx, render_abnt_pdf
from noctusai_lib.integrations.documents.formatting import FormattedDocument, ParagraphKind

from app.modules.card_hub.contrato_gerador.contexto import montar_contexto
from app.modules.card_hub.contrato_gerador.dados import DadosContrato
from app.modules.card_hub.contrato_gerador.derivacao import _hoje_padrao
from app.modules.card_hub.contrato_gerador.modelo_texto import linhas_do_template
from app.modules.card_hub.contrato_gerador.numeracao import ContadorParagrafos
from app.modules.card_hub.contrato_gerador.politica import Politica

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_PDF = "application/pdf"

#: Word paragraph-style PREFIXES that pick a non-`Normal` style while
#: building `template_bytes()` — line-level, per contract §5.2. "IMÓVEL:"
#: is the ONE template line quoting the matrícula literally (spec's
#: property-description clause); everything else defaults to `Normal`.
_TITULO = ("INSTRUMENTO PARTICULAR",)
_TITULO_SECAO = ("CLÁUSULA", "TESTEMUNHAS:")
_CITACAO = ("IMÓVEL:",)

#: The inverse mapping `gerar_pdf` reads back off the RENDERED `.docx`'s
#: paragraph styles — the two must name the same Word style strings.
_ESTILO_PARA_KIND = {
    "Title": ParagraphKind.TITLE,
    "Heading 1": ParagraphKind.HEADING,
    "Quote": ParagraphKind.QUOTE,
}


def _estilo_da_linha(linha: str) -> str:
    if linha.startswith(_TITULO):
        return "Title"
    if linha.startswith(_TITULO_SECAO):
        return "Heading 1"
    if linha.startswith(_CITACAO):
        return "Quote"
    return "Normal"


@lru_cache(maxsize=1)
def template_bytes() -> bytes:
    doc = Document()
    for linha in linhas_do_template():
        doc.add_paragraph(linha, style=_estilo_da_linha(linha))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _classificar_paragrafo(estilo: str, _texto: str) -> ParagraphKind:
    """Word paragraph style name -> ABNT `ParagraphKind` — `template_bytes()`
    is the only place that ASSIGNS the style; this only ever READS it back,
    so the rendered PDF's layout can never drift from what the template
    declared (never from the source `.docx`'s own fonts/margins)."""
    return _ESTILO_PARA_KIND.get(estilo, ParagraphKind.BODY)


def gerar_pdf(docx: bytes) -> bytes:
    """The rendered `.docx` -> ABNT PDF bytes — the SAME rendering, converted;
    both are stored (see this module's header). The PDF's metadata title is
    the already-rendered TITLE paragraph's own
    text — one source, never a second hand-built title string to drift
    from it. Raises `UnsupportedGlyphError` (from `render_abnt_pdf`) for a
    character the core Times font cannot represent; `service.gerar`
    surfaces it as a refusal, never a silent 500."""
    paragrafos = paragraphs_from_docx(docx, classify=_classificar_paragrafo)
    titulo = paragrafos[0].text if paragrafos else None
    return render_abnt_pdf(FormattedDocument(paragrafos, title=titulo))


@dataclass
class Renderizado:
    docx: bytes
    paragrafos: list[str]
    #: clause key -> how many times the template read `cl.<key>.ref`.
    referencias: dict[str, int]
    #: clause key -> its number, for the included clauses.
    clausulas: dict[str, int]


def paragrafos_do_docx(docx: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(docx)).paragraphs]


def document_xml(docx: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


def renderizar(
    adapter: DocxRenderAdapter,
    dados: DadosContrato,
    switches: dict[str, bool],
    politica: Politica,
    assinatura: date,
    hoje: Optional[date] = None,
) -> Renderizado:
    """`hoje` (default None -> `derivacao._hoje_padrao`) is the E1 empresa-
    classification reference `montar_contexto` needs — `service.gerar`
    threads ITS OWN `hoje()` snapshot through so a single generation run
    reads one "today". Resolved ONCE here (never re-defaulted per pass) so
    the two determinism-checked render passes below can never disagree
    across a real midnight boundary."""
    hoje = hoje or _hoje_padrao()
    tpl = template_bytes()
    contagem = ContadorParagrafos(None)
    adapter.render(
        tpl, montar_contexto(dados, switches, politica, assinatura, contagem, adapter, hoje)
    )

    rotulos = ContadorParagrafos(contagem.chamadas)
    contexto = montar_contexto(dados, switches, politica, assinatura, rotulos, adapter, hoje)
    docx = adapter.render(tpl, contexto)
    if rotulos.chamadas != contagem.chamadas:
        raise RuntimeError(
            "par() chamado de forma diferente nas duas passagens — template não determinístico"
        )
    registro = contexto["cl"]
    return Renderizado(
        docx=docx,
        paragrafos=paragrafos_do_docx(docx),
        referencias=dict(registro.referencias),
        clausulas={chave: c.n for chave, c in registro.items()},
    )


__all__ = [
    "MIME_DOCX",
    "MIME_PDF",
    "Renderizado",
    "document_xml",
    "gerar_pdf",
    "paragrafos_do_docx",
    "renderizar",
    "template_bytes",
]
