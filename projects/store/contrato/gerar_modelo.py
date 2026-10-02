"""Generic (data-free) compra e venda template, rendered by social-wiring's own
contract generator: same TEMPLATE wording, clause numbering, paragraph
labels, page face and ABNT PDF. Every data slot becomes a [CAMPO] blank,
highlighted in the .docx. No DadosContrato is loaded, so no real party,
imóvel or value can reach the output.
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from noctusai_lib.integrations.docx_render import apply_inline_markup, get_docx_render_adapter
from noctusai_lib.integrations.documents.formatting import Run

from app.modules.card_hub.contrato_gerador import frases
from app.modules.card_hub.contrato_gerador.documento import gerar_pdf, template_bytes
from app.modules.card_hub.contrato_gerador.numeracao import ContadorParagrafos, numerar_clausulas

OUT = Path(sys.argv[1])


def f(rotulo: str) -> str:
    """A fill-in blank."""
    return f"[{rotulo}]"


class Generico:
    """`Concordancia`'s interface with both genders merged: "VENDEDOR(A)",
    "legítimo(a) proprietário(a)". Singular — the buyer adapts to plural."""

    plural = False

    def __init__(self, base: str) -> None:
        self._base = base

    @staticmethod
    def _merge(m: str, fem: str) -> str:
        if m == fem:
            return m
        out = []
        for wm, wf in zip(m.split(" "), fem.split(" ")):
            if wm == wf:
                out.append(wm)
                continue
            i = 0
            while i < min(len(wm), len(wf)) and wm[i] == wf[i]:
                i += 1
            out.append(f"{wm}({wf[i:]})" if i == len(wm) else f"{wm}({wf[i:]})" if i else f"{wm}({wf})")
        return " ".join(out)

    def g(self, m, fem, p=None):
        return self._merge(m, fem)

    def pl(self, singular, plural):
        return singular

    NOME = property(lambda s: f"{s._base}(A)")
    ART = property(lambda s: "O(A)")
    art = property(lambda s: "o(a)")
    aos = property(lambda s: "ao(à)")
    dos = property(lambda s: "do(a)")
    pelos = property(lambda s: "pelo(a)")
    estes = property(lambda s: "este(a)")


def qualificacao(lado: str) -> str:
    return (
        f"**{f('NOME COMPLETO DO ' + lado)}**, {f('nacionalidade')}, {f('estado civil')}, {f('profissão')}, "
        f"portador(a) da cédula de identidade RG {f('nº RG / órgão emissor')} e inscrito(a) no CPF/MF {f('nº CPF')}, "
        f"com endereço eletrônico: {f('e-mail')}, residente e domiciliado(a) na {f('endereço completo com CEP')}"
    )


def brl(rotulo: str) -> str:
    return f"**R$ {f(rotulo)} ({f(rotulo.lower() + ' por extenso')})**"


def dias(rotulo) -> str:
    return f"{f('nº')} ({f('por extenso')}) dias"


#: The office's certidão list (`frases.CERTIDOES`), region-neutral: the
#: São Paulo courts/systems become the buyer's own state.
def certidoes_partes() -> list[str]:
    vistos, itens = set(), []
    for tipo, rotulo, num, _pf, _pj, _sis in frases.CERTIDOES:
        texto = (
            rotulo.replace("{R} ", "")
            .replace("do Estado de São Paulo", "do Estado de [UF]")
            .replace("ao Estado de São Paulo", "ao Estado de [UF]")
        )
        if texto in vistos:
            continue
        vistos.add(texto)
        itens.append(f"{texto} – {num} {f('____')} - emitida em {f('__/__/____')}")
    return itens


def contexto(variante: str, cont: ContadorParagrafos, adapter) -> dict:
    parcelado = variante == "parcelado"
    sw = {
        "tem_confissao": parcelado,
        "tem_declaracao_partes": True,
        "tem_assinatura_digital": False,
        "tem_intermediacao": True,
    }
    if parcelado:
        parcelas = [
            {"num": "01", "texto": f" **Sinal e princípio de pagamento:** {brl('VALOR DO SINAL')}, com vencimento em {f('data')}, por meio de {f('PIX / transferência bancária')} a ser realizada em favor do(a) **VENDEDOR(A)**: {f('NOME DO FAVORECIDO')}, {f('CPF/CNPJ')}, {f('Banco, Agência, Conta')} ou Chave PIX: {f('chave')}, operando-se automaticamente a quitação em favor do(a) Comprador(a) com o efetivo crédito na conta corrente ora indicada pelo(a) Vendedor(a)."},
            {"num": "02", "texto": f" {brl('VALOR DE CADA PARCELA')}, em {f('nº')} ({f('por extenso')}) parcelas mensais e sucessivas, vencendo-se a primeira em {f('data')} e as demais no mesmo dia dos meses subsequentes, por meio de {f('PIX / transferência bancária')} na mesma conta corrente anteriormente informada, acrescidos de {f('%')} de juros a.m., calculados pro rata die."},
        ]
        p_ref = {"sinal": "01"}
        posse_marco = "do pagamento integral do preço"
    elif variante == "avista":
        parcelas = [
            {"num": "01", "texto": f" **Sinal e princípio de pagamento:** {brl('VALOR DO SINAL')}, com vencimento em {f('data')}, por meio de {f('PIX / transferência bancária')} a ser realizada em favor do(a) **VENDEDOR(A)**: {f('NOME DO FAVORECIDO')}, {f('CPF/CNPJ')}, {f('Banco, Agência, Conta')} ou Chave PIX: {f('chave')}, operando-se automaticamente a quitação em favor do(a) Comprador(a) com o efetivo crédito na conta corrente ora indicada pelo(a) Vendedor(a)."},
            {"num": "02", "texto": f" {brl('SALDO DO PREÇO')}, {f('evento — ex.: no ato da assinatura da escritura pública de venda e compra, a ser lavrada em até 30 dias')}, por meio de {f('PIX / transferência bancária')} na mesma conta corrente anteriormente informada."},
        ]
        p_ref = {"sinal": "01"}
        posse_marco = "do recebimento da parcela 02"
    else:
        parcelas = [
            {"num": "01", "texto": f" **Sinal e princípio de pagamento:** {brl('VALOR DO SINAL')}, com vencimento em {f('data')}, por meio de {f('PIX / transferência bancária')} a ser realizada em favor do(a) **VENDEDOR(A)**: {f('NOME DO FAVORECIDO')}, {f('CPF/CNPJ')}, {f('Banco, Agência, Conta')} ou Chave PIX: {f('chave')}, operando-se automaticamente a quitação em favor do(a) Comprador(a) com o efetivo crédito na conta corrente ora indicada pelo(a) Vendedor(a)."},
            {"num": "02", "texto": f" {brl('VALOR COM RECURSOS PRÓPRIOS')}, por ocasião da assinatura do Contrato de Financiamento Imobiliário, previsto para quitação da Parcela 03, por meio de {f('PIX / transferência bancária')} na mesma conta corrente anteriormente informada."},
            {"num": "03", "texto": f" {brl('VALOR FINANCIADO')}, por meio de recursos de financiamento imobiliário e/ou moeda corrente nacional, {f('prazo / evento — ex.: no prazo de até 90 dias da assinatura do presente')}."},
        ]
        p_ref = {"sinal": "01", "financiamento": "03"}
        posse_marco = "do recebimento da parcela 03"

    return {
        **sw,
        "cl": numerar_clausulas(sw),
        "par": cont,
        "V": Generico("VENDEDOR"),
        "C": Generico("COMPRADOR"),
        "V_qualificacao": qualificacao("VENDEDOR"),
        "C_qualificacao": qualificacao("COMPRADOR"),
        "titulo_aquisitivo": f"por meio de {f('título de aquisição — ex.: escritura de compra e venda registrada sob R.__')}",
        "imovel": {
            "titulo_curto": f("IDENTIFICAÇÃO DO IMÓVEL"),
            "cidade": f("CIDADE"),
            "uf": f("UF"),
            "descricao_matricula": adapter.rich_text([Run(f("transcrever aqui a descrição do imóvel exatamente como consta na matrícula"))]),
            "inscricao_municipal": f("nº inscrição municipal / IPTU"),
            "matricula_numero": f("nº matrícula"),
            "cartorio": f("nº Oficial de Registro de Imóveis da Comarca de ______"),
            "endereco_curto": f("endereço do imóvel"),
            "em_condominio": True,
        },
        "tem_itens_integrantes": True,
        "itens_integrantes": f("listar móveis, armários, equipamentos — ou excluir este parágrafo") + ".",
        "ad_corpus": False,
        "preco": "VALOR TOTAL",
        "brl": brl,
        "dias": dias,
        "pct_extenso": lambda _x: f"{f('%')} ({f('por extenso')})",
        "parcelas": parcelas,
        "tem_financiamento": variante == "financiado",
        "p_ref": p_ref,
        "tem_permuta": False,
        "confissao": {
            "parcelas_nums": "02",
            "total": "SALDO PARCELADO",
            "juros_am": None,
            "garantia_texto": f("descrever a garantia — ex.: nota promissória, alienação fiduciária — ou excluir este parágrafo"),
        },
        "certidoes": {
            "apresentantes_texto": "O(A) **VENDEDOR(A)**",
            "apresenta": "apresenta",
            "seus_nomes": "seu nome",
            "grupos": [{"num": "1", "em_nome_de": f("NOME DO VENDEDOR"), "sufixo": None, "itens": certidoes_partes()}],
            "grupos_imovel": [{"num": "2", "titulo": f("IDENTIFICAÇÃO DO IMÓVEL"), "itens": [
                f"Visualização da Certidão da Matrícula do imóvel nº {f('nº matrícula')}, emitida em {f('__/__/____')}",
                f"Certidão Negativa de Débitos Municipais (IPTU) nº {f('____')} - emitida em {f('__/__/____')}",
                f"Certidão Negativa de Débitos Condominiais - emitida em {f('__/__/____')}",
            ]}],
            "pendencias": [
                {"letra": l, "texto": t} for l, t in zip("abcdef", [
                    frases.PENDENCIA_DOCUMENTOS,
                    "Certidão de estado civil atualizada (nascimento ou casamento), emitida há menos de 90 dias",
                    frases.PENDENCIA_CONTAS_CONSUMO,
                    frases.PENDENCIA_CONDOMINIO,
                    frases.PENDENCIA_MATRICULA,
                ])
            ],
        },
        "prazo_esclarecimentos": f("nº"),
        "prazo_pendencias": "PRAZO",
        "tem_saldo_devedor": False,
        "onus": {},
        "posse": {
            "prazo": "PRAZO",
            "marco_texto": posse_marco,
            "condicao_frase": "",
            "multa_diaria": "MULTA DIÁRIA",
        },
        "tem_multa_diaria_posse": True,
        "rescisao": {"cura_frase": ""},
        "multa_rescisoria": "MULTA RESCISÓRIA",
        "assinatura": {"local": f("Cidade"), "data_extenso": f("dia") + " de " + f("mês") + " de " + f("ano")},
        "resolutiva_notificacao_email": False,
        "corretagem": {
            "contratantes_texto": "o(a) **VENDEDOR(A)**",
            "contratantes_texto_cap": "O(A) **VENDEDOR(A)**",
            "contrata": "contrata",
            "empresas_texto": "o(a) intermediador(a) abaixo qualificado(a)",
            "contratadas_texto": "o(a) **INTERMEDIADOR(A)**",
            "qualificados": [f"**{f('NOME / RAZÃO SOCIAL DA IMOBILIÁRIA OU CORRETOR(A)')}**, inscrito(a) no {f('CPF/CNPJ')}, com inscrição no CRECI sob o nº {f('CRECI')}, endereço eletrônico: {f('e-mail')}, com sede na {f('endereço completo')}."],
            "total": "VALOR DA CORRETAGEM",
            "parcelamento_texto": "em uma única parcela",
            "marcos_texto": "da Parcela 01",
            "splits": [f"{brl('VALOR')}, por meio de depósito bancário em favor de {f('NOME')}, {f('CPF/CNPJ')}, {f('Banco, Agência, Conta')} ou Chave PIX: {f('chave')}"],
            "pct_rescisao": f("%"),
        },
        "foro": {"comarca": f("COMARCA")},
        "vias": "02 (duas)",
        "linha_assinatura": frases.LINHA_ASSINATURA,
        "V_assinantes_fisicos": [{"nome": f("NOME DO(A) VENDEDOR(A)"), "documento": f"CPF {f('nº')}"}],
        "C_assinantes_fisicos": [{"nome": f("NOME DO(A) COMPRADOR(A)"), "documento": f"CPF {f('nº')}"}],
        "testemunhas_fisicas": [{"nome": f("NOME DA TESTEMUNHA 1"), "documento": f"CPF {f('nº')}"},
                                 {"nome": f("NOME DA TESTEMUNHA 2"), "documento": f"CPF {f('nº')}"}],
    }


BLANK = re.compile(r"(\[[^\]]+\])")


def destacar_lacunas(docx: bytes) -> bytes:
    """Yellow-highlight every [CAMPO] so the buyer sees what to fill."""
    doc = Document(io.BytesIO(docx))
    for p in doc.paragraphs:
        for run in list(p.runs):
            partes = BLANK.split(run.text)
            if len(partes) == 1:
                continue
            run.text = partes[0]
            ancora = run
            for parte in partes[1:]:
                if not parte:
                    continue
                novo = p.add_run(parte)
                novo.bold, novo.underline, novo.italic = run.bold, run.underline, run.italic
                novo.font.name, novo.font.size = run.font.name, run.font.size
                if BLANK.fullmatch(parte):
                    novo.font.highlight_color = WD_COLOR_INDEX.YELLOW
                ancora._r.addnext(novo._r)
                ancora = novo
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    adapter = get_docx_render_adapter()
    tpl = template_bytes()
    OUT.mkdir(parents=True, exist_ok=True)
    for variante, nome in (("avista", "Contrato-Compra-e-Venda-Imovel-A-Vista"),
                           ("financiado", "Contrato-Compra-e-Venda-Imovel-Financiado"),
                           ("parcelado", "Contrato-Compra-e-Venda-Imovel-Parcelado-Direto")):
        conta = ContadorParagrafos(None)
        adapter.render(tpl, contexto(variante, conta, adapter))
        rotulos = ContadorParagrafos(conta.chamadas)
        docx = apply_inline_markup(adapter.render(tpl, contexto(variante, rotulos, adapter)))
        (OUT / f"{nome}.pdf").write_bytes(gerar_pdf(docx))
        (OUT / f"{nome}.docx").write_bytes(destacar_lacunas(docx))
        print("ok", nome)


if __name__ == "__main__":
    main()
