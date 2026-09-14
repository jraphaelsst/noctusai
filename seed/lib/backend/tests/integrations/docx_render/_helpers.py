"""Programmatic `.docx` template builders shared by this directory's tests.

No binary fixtures, no real contract content — every template is built
in-test via `python-docx`. NOT a `test_*.py` module (pytest won't collect
it as a test file); imported by the actual test modules.
"""

from __future__ import annotations

import io

from docx import Document


def simple_placeholder_docx(text: str) -> bytes:
    """A one-paragraph template: `text` verbatim (embed `{{ ... }}` in it)."""
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def conditional_and_loop_docx() -> bytes:
    """A template exercising every shape the S-c brief calls out:

    - a plain `{{ var }}` placeholder
    - a `{% if %}` conditional block
    - a `{% for %}` table-row loop (docxtpl's `{%tr for %}` / `{%tr endfor %}`
      row-repeat convention — the for/endfor tags each occupy their OWN
      dedicated table row; the row(s) between them repeat once per item)
    """
    doc = Document()
    doc.add_paragraph("Contratante: {{ nome_contratante }}")
    doc.add_paragraph("{% if tem_fiador %}Fiador: {{ nome_fiador }}{% endif %}")

    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "{%tr for p in parcelas %}"
    table.rows[0].cells[1].text = ""
    table.rows[1].cells[0].text = "{{ p.numero }}"
    table.rows[1].cells[1].text = "{{ p.valor }}"
    table.rows[2].cells[0].text = "{%tr endfor %}"
    table.rows[2].cells[1].text = ""

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def read_paragraph_text(docx_bytes: bytes, index: int = 0) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    return doc.paragraphs[index].text


def read_all_paragraph_text(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs]


def read_table_rows(docx_bytes: bytes, table_index: int = 0) -> list[list[str]]:
    doc = Document(io.BytesIO(docx_bytes))
    table = doc.tables[table_index]
    return [[cell.text for cell in row.cells] for row in table.rows]
