"""`apply_inline_markup` — the matrícula transcriber's `**`/`<u>` markup model
turned into real runs in a rendered `.docx`. Synthetic documents only.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document

from noctusai_lib.integrations.documents.formatting import Run
from noctusai_lib.integrations.docx_render import (
    DocxtplRenderAdapter,
    FakeDocxRenderAdapter,
    apply_inline_markup,
    paragraphs_with_raw_markup,
)

from ._helpers import simple_placeholder_docx


def _runs(docx_bytes: bytes, index: int = 0) -> list[tuple[str, bool, bool]]:
    paragraph = Document(io.BytesIO(docx_bytes)).paragraphs[index]
    return [(r.text, bool(r.bold), bool(r.underline)) for r in paragraph.runs]


def _xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


class TestSplitsMarkedRuns:
    def test_bold_and_underline_become_runs_and_the_markers_disappear(self) -> None:
        doc = simple_placeholder_docx("O **VENDEDOR** vende a <u>CLÁUSULA PRIMEIRA</u> hoje.")
        out = apply_inline_markup(doc)
        assert _runs(out) == [
            ("O ", False, False),
            ("VENDEDOR", True, False),
            (" vende a ", False, False),
            ("CLÁUSULA PRIMEIRA", False, True),
            (" hoje.", False, False),
        ]
        assert "**" not in _xml(out)
        assert "&lt;u&gt;" not in _xml(out)
        assert paragraphs_with_raw_markup(out) == []

    def test_nested_markers_compose(self) -> None:
        out = apply_inline_markup(simple_placeholder_docx("a **b <u>c</u>** d"))
        assert _runs(out) == [
            ("a ", False, False),
            ("b ", True, False),
            ("c", True, True),
            (" d", False, False),
        ]

    def test_a_line_break_inside_a_marked_run_survives(self) -> None:
        out = apply_inline_markup(simple_placeholder_docx("**linha um\nlinha dois**"))
        assert Document(io.BytesIO(out)).paragraphs[0].text == "linha um\nlinha dois"
        assert all(bold for _t, bold, _u in _runs(out))

    def test_other_run_properties_are_copied_onto_each_piece(self) -> None:
        doc = Document()
        run = doc.add_paragraph().add_run("x **y** z")
        run.italic = True
        run.font.name = "Arial"
        buf = io.BytesIO()
        doc.save(buf)
        out = Document(io.BytesIO(apply_inline_markup(buf.getvalue())))
        pieces = out.paragraphs[0].runs
        assert [p.text for p in pieces] == ["x ", "y", " z"]
        assert all(p.italic and p.font.name == "Arial" for p in pieces)
        assert [bool(p.bold) for p in pieces] == [False, True, False]

    def test_table_cells_are_converted_too(self) -> None:
        doc = Document()
        doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "**negrito**"
        buf = io.BytesIO()
        doc.save(buf)
        out = Document(io.BytesIO(apply_inline_markup(buf.getvalue())))
        cell_runs = out.tables[0].rows[0].cells[0].paragraphs[0].runs
        assert [(r.text, bool(r.bold)) for r in cell_runs] == [("negrito", True)]


class TestComposesWithExistingFormatting:
    def test_rich_text_runs_keep_their_captured_formatting(self) -> None:
        """The matrícula quote shape: a docxtpl `{{r }}` slot's runs (already
        bold/underlined) next to marked template text — both survive."""
        adapter = DocxtplRenderAdapter()
        template = simple_placeholder_docx("**IMÓVEL:** {{r quote }} sob nº **{{ n }}**.")
        quote = adapter.rich_text((Run("apto 11 de "), Run("FULANO", bold=True), Run(" R.1", underline=True)))
        out = apply_inline_markup(adapter.render(template, {"quote": quote, "n": "123"}))
        runs = [r for r in _runs(out) if r[0]]
        assert ("IMÓVEL:", True, False) in runs
        assert ("FULANO", True, False) in runs
        assert (" R.1", False, True) in runs
        assert ("123", True, False) in runs
        assert Document(io.BytesIO(out)).paragraphs[0].text == "IMÓVEL: apto 11 de FULANO R.1 sob nº 123."

    def test_a_run_already_bold_stays_bold_where_unmarked(self) -> None:
        doc = Document()
        run = doc.add_paragraph().add_run("tudo <u>negrito</u>")
        run.bold = True
        buf = io.BytesIO()
        doc.save(buf)
        assert _runs(apply_inline_markup(buf.getvalue())) == [
            ("tudo ", True, False),
            ("negrito", True, True),
        ]


class TestNeverGuesses:
    def test_an_unbalanced_marker_is_kept_literal_and_reported(self) -> None:
        out = apply_inline_markup(simple_placeholder_docx("sem par ** aqui"))
        assert Document(io.BytesIO(out)).paragraphs[0].text == "sem par ** aqui"
        assert paragraphs_with_raw_markup(out) == [0]

    def test_a_document_without_markers_is_textually_unchanged(self) -> None:
        doc = simple_placeholder_docx("Texto simples, sem ênfase.")
        out = apply_inline_markup(doc)
        assert _runs(out) == [("Texto simples, sem ênfase.", False, False)]


class TestDeterminismAndAdapters:
    def test_output_is_byte_deterministic(self) -> None:
        doc = simple_placeholder_docx("**a** b <u>c</u>")
        assert apply_inline_markup(doc) == apply_inline_markup(doc)

    def test_runs_on_the_fake_adapters_output(self) -> None:
        fake = FakeDocxRenderAdapter()
        out = apply_inline_markup(fake.render(simple_placeholder_docx("{{ x }}"), {"x": "1"}))
        assert Document(io.BytesIO(out)).paragraphs[0].text.startswith("FAKE-DOCX-RENDER")
