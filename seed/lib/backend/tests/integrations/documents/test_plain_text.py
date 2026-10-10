import asyncio
import io

import docx
import fitz

from noctusai_lib.integrations.documents.plain_text import extract_plain_text
from noctusai_lib.integrations.documents.transcription import FakeDocumentTranscriber


def run(content, name, **kw):
    return asyncio.run(extract_plain_text(content, name, **kw))


def _docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("Olá mundo")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "a"
    t.rows[0].cells[1].text = "b"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    d = fitz.open()
    d.new_page().insert_text((72, 72), "texto da camada " * 20)
    return d.tobytes()


def test_docx_paragraphs_and_tables():
    r = run(_docx(), "x.DOCX")
    assert r.error is None and r.text == "Olá mundo\na | b"


def test_pdf_via_injected_transcriber():
    r = run(_pdf(), "x.pdf", transcriber=FakeDocumentTranscriber())
    assert r.error is None and r.text


def test_text_utf8_and_cp1252():
    assert run("ação".encode("utf-8"), "a.txt").text == "ação"
    assert run("ação".encode("cp1252"), "a.md").text == "ação"


def test_encoding_error():
    assert run(b"\x81\x8d\x8f\x90\x9d", "a.csv").error == "encoding"


def test_magic_mismatch():
    assert run(b"not a pdf", "a.pdf").error == "formato_invalido"
    assert run(b"not a zip", "a.docx").error == "formato_invalido"
    assert run(_pdf(), "a.txt").error == "formato_invalido"
    assert run(_docx(), "a.txt").error == "formato_invalido"
    assert run(b"x", "a.exe").error == "formato_invalido"


def test_docx_zip_without_document_xml():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("other.xml", "<a/>")
    assert run(buf.getvalue(), "a.docx").error == "formato_invalido"


def test_empty():
    assert run(b"", "a.txt").error == "vazio"
