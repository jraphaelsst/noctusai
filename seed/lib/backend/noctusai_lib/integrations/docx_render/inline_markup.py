"""Inline markup (`**bold**`, `<u>underline</u>`) in a rendered `.docx` → real runs.

WHY THIS EXISTS
---------------
The matrícula transcriber already has ONE inline-formatting model: the vision
reply marks bold with `**…**` and underline with `<u>…</u>`, and
`documents.transcription.parse_markup` turns that into plain text +
`FormatRange`s. A document GENERATOR that wants to emphasise its own data
(a party's name, a price, a clause label) needs exactly the same two
properties — so it writes the same markers into its wording, renders the
template as usual, and hands the result to `apply_inline_markup`, which runs
the SAME `parse_markup` over every run and splits it into formatted runs.
One markup model end to end: nothing about emphasis is re-invented here.

Contract:

- Works on ANY `.docx` bytes (a `DocxRenderAdapter.render()` result — Real or
  Fake — or a hand-built document); pure, deterministic, no vendor SDK
  (python-docx is a core seed dependency), so there is no Fake/Real split.
- Run-local: a marker pair must open and close inside ONE run. Text a
  template wrote on one line is one run after rendering; a docxtpl
  `{{r … }}` rich-text slot is its own run(s) — so never wrap a rich slot in
  markers. A pair that straddles runs is unbalanced within each run, and
  `parse_markup` keeps an unbalanced marker LITERAL (and logs) rather than
  guessing — `has_raw_markup` over the result then finds it, which is how a
  consumer turns that into a loud refusal (see `paragraphs_with_raw_markup`).
- Composes with formatting the run already had: a run that was bold stays
  bold everywhere, a marked stretch ADDS bold/underline — so a matrícula
  quote's captured runs survive untouched (they carry no markers: a text
  with markers left in it is refused upstream, `has_raw_markup`).
- Every other run property (font, size, colour, style) is copied from the
  original run onto each piece, so the split is invisible except for the
  emphasis it adds.
"""
from __future__ import annotations

import io
from copy import deepcopy
from typing import Iterator

from docx import Document
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run as DocxRun

from noctusai_lib.integrations.documents.abnt import runs_from_ranges
from noctusai_lib.integrations.documents.transcription import has_raw_markup, parse_markup
from noctusai_lib.integrations.docx_render._zip import without_wall_clock


def _paragraphs(document) -> Iterator[DocxParagraph]:
    """Body paragraphs, then every table cell's (nested tables included)."""
    yield from document.paragraphs

    def _tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    yield from _tables(cell.tables)

    yield from _tables(document.tables)


def _split_run(run: DocxRun, paragraph: DocxParagraph) -> None:
    texto, ranges = parse_markup(run.text)
    pedacos = runs_from_ranges(texto, ranges)
    anchor = run._r
    for pedaco in pedacos:
        novo = deepcopy(anchor)
        anchor.addprevious(novo)
        alvo = DocxRun(novo, paragraph)
        alvo.text = pedaco.text  # python-docx maps "\n" → <w:br/>, "\t" → <w:tab/>
        if pedaco.bold:
            alvo.bold = True
        if pedaco.underline:
            alvo.underline = True
    anchor.getparent().remove(anchor)


def apply_inline_markup(docx_bytes: bytes) -> bytes:
    """Return `docx_bytes` with every `**…**` / `<u>…</u>` inside a run
    replaced by formatted runs (markers removed). A document with no marker
    comes back re-saved but textually identical."""
    document = Document(io.BytesIO(docx_bytes))
    for paragraph in _paragraphs(document):
        for run in list(paragraph.runs):
            if has_raw_markup(run.text):
                _split_run(run, paragraph)
    out = io.BytesIO()
    document.save(out)
    return without_wall_clock(out.getvalue())


def paragraphs_with_raw_markup(docx_bytes: bytes) -> list[int]:
    """0-based indexes of body paragraphs whose text still carries a literal
    marker — what a caller checks AFTER `apply_inline_markup` to refuse a
    document with an unbalanced/straddling pair instead of shipping `**`."""
    document = Document(io.BytesIO(docx_bytes))
    return [i for i, p in enumerate(document.paragraphs) if has_raw_markup(p.text)]


__all__ = ["apply_inline_markup", "paragraphs_with_raw_markup"]
