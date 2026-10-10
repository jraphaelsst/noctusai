"""Uploaded file -> plain text (PDF / DOCX / TXT / MD / CSV). Never raises on bad input.

`extract_plain_text` is ASYNC because the PDF path goes through the seed
`DocumentTranscriber` (text layer first, vision per page), which is async.
DOCX and text formats are decoded synchronously inside it.

Error codes (`PlainText.error`): `formato_invalido` (extension unsupported or
magic bytes disagree with it), `encoding` (text file neither UTF-8 nor
cp1252), `vazio` (nothing extractable), `docx_ilegivel` (corrupt DOCX), or the
transcriber's own code for a PDF it could not read.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from typing import Optional

from noctusai_lib.integrations.documents.transcription import DocumentTranscriber

_TEXT_EXTS = frozenset({"txt", "md", "csv"})
_SUPPORTED = _TEXT_EXTS | {"pdf", "docx"}


@dataclass(frozen=True)
class PlainText:
    text: str = ""
    error: Optional[str] = None

    @property
    def ok(self) -> bool:
        return self.error is None


def docx_to_text(content: bytes) -> str:
    """DOCX bytes -> paragraphs, then table rows (cells joined by ' | '). Raises on a corrupt file."""
    import docx  # python-docx (declared seed dependency)

    doc = docx.Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def _is_docx_zip(content: bytes) -> bool:
    if not content.startswith(b"PK\x03\x04"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            return "word/document.xml" in zf.namelist()
    except zipfile.BadZipFile:
        return False


def _decode(content: bytes) -> Optional[str]:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return content.decode("cp1252")
    except UnicodeDecodeError:
        return None


async def extract_plain_text(
    content: bytes,
    filename: str,
    *,
    transcriber: Optional[DocumentTranscriber] = None,
) -> PlainText:
    """Extract text from `content`. `transcriber` defaults to the real ladder transcriber."""
    ext = _ext(filename)
    if ext not in _SUPPORTED or not content:
        return PlainText(error="formato_invalido" if ext not in _SUPPORTED else "vazio")

    if ext == "pdf":
        if not content.startswith(b"%PDF"):
            return PlainText(error="formato_invalido")
        if transcriber is None:
            from noctusai_lib.integrations.documents.transcription import (
                make_document_transcriber,
            )

            transcriber = make_document_transcriber(real=True)
        result = await transcriber.transcribe(
            content, mimetype="application/pdf", filename=filename
        )
        if not result.ok:
            return PlainText(error=result.error)
        return PlainText(text=result.text) if result.text.strip() else PlainText(error="vazio")

    if ext == "docx":
        if not _is_docx_zip(content):
            return PlainText(error="formato_invalido")
        try:
            text = docx_to_text(content)
        except Exception:  # noqa: BLE001 - corrupt package: typed error, not a raise
            return PlainText(error="docx_ilegivel")
        return PlainText(text=text) if text.strip() else PlainText(error="vazio")

    # txt / md / csv: must not be a binary container pretending to be text.
    if content.startswith((b"%PDF", b"PK\x03\x04")) or b"\x00" in content:
        return PlainText(error="formato_invalido")
    text = _decode(content)
    if text is None:
        return PlainText(error="encoding")
    return PlainText(text=text) if text.strip() else PlainText(error="vazio")
